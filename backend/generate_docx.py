import sys
import json
import os
import traceback
import subprocess
import platform
import tempfile
import uuid
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, ns
from docx.oxml.ns import qn
import markdown2
from bs4 import BeautifulSoup, NavigableString
from PIL import Image

# Force UTF-8
if hasattr(sys.stdin, 'reconfigure'):
    sys.stdin.reconfigure(encoding='utf-8')
    sys.stdout.reconfigure(encoding='utf-8')

# --- XML UTILITIES ---
def create_element(name): return OxmlElement(name)
def create_attribute(element, name, value): element.set(ns.qn(name), value)

def hex_to_rgb(hex_str):
    """Converts #RRGGBB to docx RGBColor."""
    hex_str = str(hex_str).strip().lstrip('#')
    if len(hex_str) != 6: return RGBColor(0, 0, 0)
    try:
        return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))
    except:
        return RGBColor(0, 0, 0)

def set_document_background(doc, hex_color):
    """Injects a full-page background color via XML if not white."""
    if not hex_color or hex_color.lower() in ['#ffffff', 'white', '']: return
    hex_color = hex_color.lstrip('#')
    
    background = OxmlElement('w:background')
    background.set(qn('w:color'), hex_color)
    doc._element.insert(0, background)
    
    # Force Word to display the background
    settings = doc.settings.element
    display_bg = OxmlElement('w:displayBackgroundShape')
    settings.append(display_bg)

def add_floating_image(run, image_path, width_in_emu, height_in_emu, pos_x, pos_y, behind_text=False):
    """Injects a floating image using absolute positioning on the page."""
    try:
        picture = run.add_picture(image_path)
        drawing = picture._inline.getparent()
        
        anchor = OxmlElement('wp:anchor')
        anchor.set('distT', "0")
        anchor.set('distB', "0")
        anchor.set('distL', "0")
        anchor.set('distR', "0")
        anchor.set('simplePos', "0")
        anchor.set('relativeHeight', "251658240")
        anchor.set('behindDoc', "1" if behind_text else "0") 
        anchor.set('locked', "0")
        anchor.set('layoutInCell', "1")
        anchor.set('allowOverlap', "1")
        
        simplePos = OxmlElement('wp:simplePos')
        simplePos.set('x', "0")
        simplePos.set('y', "0")
        anchor.append(simplePos)
        
        positionH = OxmlElement('wp:positionH')
        positionH.set('relativeFrom', "page")
        posOffsetH = OxmlElement('wp:posOffset')
        posOffsetH.text = str(int(pos_x))
        positionH.append(posOffsetH)
        anchor.append(positionH)
        
        positionV = OxmlElement('wp:positionV')
        positionV.set('relativeFrom', "page")
        posOffsetV = OxmlElement('wp:posOffset')
        posOffsetV.text = str(int(pos_y))
        positionV.append(posOffsetV)
        anchor.append(positionV)
        
        extent = OxmlElement('wp:extent')
        extent.set('cx', str(int(width_in_emu)))
        extent.set('cy', str(int(height_in_emu)))
        anchor.append(extent)
        
        anchor.append(OxmlElement('wp:wrapNone'))
        anchor.append(picture._inline.docPr)
        
        graphic = picture._inline.graphic
        xfrm = graphic.graphicData.xpath('.//pic:spPr/a:xfrm')[0]
        ext = xfrm.xpath('./a:ext')[0]
        ext.set('cx', str(int(width_in_emu)))
        ext.set('cy', str(int(height_in_emu)))
        
        anchor.append(graphic)
        drawing.remove(picture._inline)
        drawing.append(anchor)
    except Exception as e:
        print(f"Image positioning warning: {e}", file=sys.stderr)

def add_horizontal_line(doc, size='4', color='000000', full_bleed=False, left_margin_pt=72, right_margin_pt=72):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if full_bleed:
        p.paragraph_format.left_indent = Pt(-left_margin_pt)
        p.paragraph_format.right_indent = Pt(-right_margin_pt)
    else:
        p.paragraph_format.left_indent = Pt(0)
        p.paragraph_format.right_indent = Pt(0)
    
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    p_pr = p._element.get_or_add_pPr()
    border = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size)) 
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), color)
    border.append(bottom)
    p_pr.append(border)

def set_font(run, font_name, size, bold=False, italic=False, color=None):
    """Enforces specific user-selected font styles directly to the run."""
    run.font.name = font_name
    run.font.size = Pt(int(size))
    run.bold = bold
    run.italic = italic
    if color: run.font.color.rgb = color
    
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def process_text_runs(paragraph, bs_element, font_name, base_font_size, base_text_color):
    if isinstance(bs_element, NavigableString):
        if bs_element.strip():
            run = paragraph.add_run(str(bs_element))
            set_font(run, font_name, base_font_size, color=base_text_color)
        return

    for child in getattr(bs_element, 'children', []):
        if child.name == 'br':
            paragraph.add_run().add_break()
        elif child.name in ['strong', 'b']:
            run = paragraph.add_run(child.get_text())
            set_font(run, font_name, base_font_size, bold=True, color=base_text_color)
        elif child.name in ['em', 'i']:
            run = paragraph.add_run(child.get_text())
            set_font(run, font_name, base_font_size, italic=True, color=base_text_color)
        elif child.name == 'a':
            run = paragraph.add_run(child.get_text())
            set_font(run, font_name, base_font_size, color=RGBColor(37, 99, 235))
        else:
            if child.string and child.string.strip():
                run = paragraph.add_run(child.string)
                set_font(run, font_name, base_font_size, color=base_text_color)

def format_markdown_to_docx(doc, md_text, font_name, base_font_size, base_text_color):
    if not md_text: return
    html = markdown2.markdown(md_text, extras=["tables", "breaks"])
    soup = BeautifulSoup(html, 'html.parser')
    
    for element in soup:
        if isinstance(element, NavigableString):
            if element.strip():
                p = doc.add_paragraph()
                run = p.add_run(str(element).strip())
                set_font(run, font_name, base_font_size, color=base_text_color)
            continue

        if element.name in ['h1', 'h2']:
            p = doc.add_paragraph()
            run = p.add_run(element.get_text())
            set_font(run, font_name, int(base_font_size) + 4, bold=True, color=base_text_color)
        elif element.name == 'h3':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(element.get_text())
            set_font(run, font_name, int(base_font_size) + 1, bold=True, color=base_text_color)
        elif element.name in ['p', 'div', 'pre', 'blockquote']:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(6)
            process_text_runs(p, element, font_name, base_font_size, base_text_color)
        elif element.name in ['ul', 'ol']:
            list_style = 'List Bullet' if element.name == 'ul' else 'List Number'
            for li in element.find_all('li'):
                p = doc.add_paragraph(style=list_style)
                p.paragraph_format.left_indent = Inches(0.4)
                process_text_runs(p, li, font_name, base_font_size, base_text_color)
        elif element.name == 'table':
            rows = element.find_all('tr')
            if rows:
                col_count = max(len(row.find_all(['th', 'td'])) for row in rows)
                table = doc.add_table(rows=0, cols=col_count)
                table.style = 'Table Grid'
                for row in rows:
                    cells = row.find_all(['th', 'td'])
                    row_cells = table.add_row().cells
                    for i, cell in enumerate(cells):
                        if i < col_count:
                            cell_text = cell.get_text(strip=True)
                            p = row_cells[i].paragraphs[0]
                            run = p.add_run(cell_text)
                            set_font(run, font_name, base_font_size, bold=(cell.name == 'th'), color=base_text_color)
            doc.add_paragraph()

# --- CRITICAL FIX: CROSS-PLATFORM LINUX PDF CONVERSION ---
def convert_to_pdf_linux(docx_path, output_dir):
    """
    Converts DOCX to PDF using LibreOffice headless mode.
    Requires ZERO Microsoft Word dependencies. Works natively on Ubuntu/Linux VPS servers.
    Includes explicit pathlib URI parsing to prevent Windows bootstrap.ini crashes.
    """
    try:
        # 1. Determine local OS executable
        if platform.system() == 'Windows':
            path1 = r"C:\Program Files\LibreOffice\program\soffice.exe"
            path2 = r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"
            
            if os.path.exists(path1):
                libreoffice_cmd = path1
            elif os.path.exists(path2):
                libreoffice_cmd = path2
            else:
                libreoffice_cmd = "soffice" # Fallback if added to PATH
        elif platform.system() == 'Darwin':
            libreoffice_cmd = "/Applications/LibreOffice.app/Contents/MacOS/soffice"
        else:
            libreoffice_cmd = "libreoffice"

        # 2. Production Hardening: Create isolated temp profile to prevent LibreOffice hanging
        # CRITICAL FIX: Use pathlib to generate a valid cross-platform file URI 
        # Prevents Windows backslashes from breaking LibreOffice VFS resulting in fake "bootstrap.ini corrupt" errors.
        profile_dir = os.path.join(tempfile.gettempdir(), f"lo_{uuid.uuid4().hex}")
        profile_uri = Path(profile_dir).as_uri()
        env_profile = f"-env:UserInstallation={profile_uri}"

        # 3. Execute Headless Conversion
        process = subprocess.run([
            libreoffice_cmd,
            env_profile,
            '--headless',
            '--convert-to', 'pdf',
            '--outdir', output_dir,
            docx_path
        ], stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
        
        # 4. Trigger safe Exit Code 2 to notify server.js of failure
        if process.returncode != 0:
            print(f"LIBREOFFICE_ERROR: {process.stderr}", file=sys.stderr)
            sys.exit(2)
            
    except Exception as e:
        print(f"PDF_CONVERSION_ERROR: {e}", file=sys.stderr)
        sys.exit(2)

# --- MAIN GENERATOR ---
def generate_document():
    try:
        file_id = sys.argv[1] if len(sys.argv) > 1 else "Final_Document"
        script_dir = os.path.dirname(os.path.abspath(__file__))
        assets_dir = os.path.join(script_dir, "assets")
        
        input_data = sys.stdin.read()
        if not input_data: raise ValueError("No input data received.")
        data = json.loads(input_data)
        
        format_type = data.get('format', 'docx')
        doc_title = data.get('documentTitle', 'Untitled Document')
        sections = data.get('sections', [])
        cust = data.get('customization', {})
        
        font_name = cust.get('font', 'Arial')
        font_size = int(cust.get('size', 12))
        header_text = cust.get('header', '')
        footer_text = cust.get('footer', '')
        
        mt = float(cust.get('marginTop', 72))
        mb = float(cust.get('marginBottom', 72))
        ml = float(cust.get('marginLeft', 72))
        mr = float(cust.get('marginRight', 72))
        
        bg_color_hex = cust.get('bgColor', '#ffffff')
        text_color_hex = cust.get('textColor', '#000000')
        text_color_rgb = hex_to_rgb(text_color_hex)
        
        # 1. MARKDOWN EXPORT
        if format_type == 'md':
            output_filename = os.path.join(script_dir, f"{file_id}.md")
            with open(output_filename, "w", encoding="utf-8") as f:
                f.write(f"<style>\nbody {{\n  background-image: url('assets/watermark.png');\n  background-repeat: repeat-y;\n  background-position: center;\n  background-size: 60%;\n")
                f.write(f"  font-family: '{font_name}', sans-serif;\n")
                f.write(f"  background-color: {bg_color_hex};\n")
                f.write(f"  color: {text_color_hex};\n")
                f.write(f"}}\n</style>\n\n")
                
                if header_text: f.write(f"<div align='right'><b>{header_text}</b></div>\n")
                f.write(f"<div align='center'>\n")
                f.write(f"<b>OFFICE NO. 12, 2ND FLOOR, NO.40 DLF INDUSTRIAL AREA, KIRTI NAGAR, DELHI 110015.</b><br>\n")
                f.write(f"<b>Contact No. : <a href='tel:+919987159986'>+91 99871 59986</a> | Email : <a href='mailto:info@sgwebapp.com'>info@sgwebapp.com</a> | Website : <a href='http://www.sgwebapp.com'>www.sgwebapp.com</a></b>\n")
                f.write(f"</div>\n<hr>\n\n")
                
                f.write(f"<h1 align='center'>{doc_title.upper()}</h1>\n\n")
                for i, sec in enumerate(sections):
                    f.write(f"## {i+1}. {sec.get('heading', 'Section')}\n\n{sec.get('content', '')}\n\n")
                
                if footer_text: f.write(f"<hr>\n<div align='center'><i>{footer_text.replace('[PAGE]', '1')}</i></div>\n")
            print("SUCCESS_MD")
            return

        # 2. MASTER DOCX GENERATION (100% UNTOUCHED)
        doc = Document()
        doc.settings.odd_and_even_pages_header_footer = False
        set_document_background(doc, bg_color_hex)
        
        for section in doc.sections:
            section.top_margin = Pt(mt)
            section.bottom_margin = Pt(mb)
            section.left_margin = Pt(ml)
            section.right_margin = Pt(mr)
            section.different_first_page_header_footer = False
            
            header = section.header
            h_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
            
            if header_text:
                h_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                h_run = h_para.add_run(header_text)
                set_font(h_run, font_name, font_size - 2, color=RGBColor(100, 100, 100))
                
            watermark_path = os.path.join(assets_dir, "watermark.png")
            if os.path.exists(watermark_path):
                w_run = h_para.add_run()
                add_floating_image(w_run, watermark_path, 6 * 914400, 6 * 914400, 1.25 * 914400, 3.5 * 914400, behind_text=True)

            footer = section.footer
            footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
            footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            if footer_text:
                if '[PAGE]' in footer_text:
                    parts = footer_text.split('[PAGE]')
                    f_run = footer_para.add_run(parts[0])
                    set_font(f_run, font_name, font_size - 2)
                    
                    fldChar1 = create_element('w:fldChar')
                    create_attribute(fldChar1, 'w:fldCharType', 'begin')
                    instrText = create_element('w:instrText')
                    create_attribute(instrText, 'xml:space', 'preserve')
                    instrText.text = "PAGE"
                    fldChar2 = create_element('w:fldChar')
                    create_attribute(fldChar2, 'w:fldCharType', 'end')
                    
                    run_pg = footer_para.add_run()
                    run_pg._r.append(fldChar1)
                    run_pg._r.append(instrText)
                    run_pg._r.append(fldChar2)
                    set_font(run_pg, font_name, font_size - 2)
                    
                    if len(parts) > 1:
                        f_run_end = footer_para.add_run(parts[1])
                        set_font(f_run_end, font_name, font_size - 2)
                else:
                    f_run = footer_para.add_run(footer_text)
                    set_font(f_run, font_name, font_size - 2)
            else:
                run_p = footer_para.add_run()
                set_font(run_p, font_name, 10)
                fldChar1 = create_element('w:fldChar')
                create_attribute(fldChar1, 'w:fldCharType', 'begin')
                instrText = create_element('w:instrText')
                create_attribute(instrText, 'xml:space', 'preserve')
                instrText.text = "PAGE"
                fldChar2 = create_element('w:fldChar')
                create_attribute(fldChar2, 'w:fldCharType', 'end')
                run_p._r.append(fldChar1)
                run_p._r.append(instrText)
                run_p._r.append(fldChar2)

        top_path = os.path.join(assets_dir, "top.png")
        if os.path.exists(top_path):
            try:
                with Image.open(top_path) as img:
                    img_w, img_h = img.size
                    aspect_ratio = img_h / img_w
            except Exception:
                aspect_ratio = 0.15 
                
            img_height_in = 8.5 * aspect_ratio
            img_height_emu = int(img_height_in * 914400)
            
            banner_para = doc.add_paragraph()
            banner_para.paragraph_format.space_before = Pt(0)
            banner_para.paragraph_format.space_after = Pt(0)
            run = banner_para.add_run()
            add_floating_image(run, top_path, 8.5 * 914400, img_height_emu, 0, 0, behind_text=True)

            top_margin_pt = mt
            img_height_pt = img_height_in * 72
            space_pt = max(0, img_height_pt - top_margin_pt + 4)
            
            address_para = doc.add_paragraph()
            address_para.paragraph_format.space_before = Pt(space_pt)
            address_para.paragraph_format.space_after = Pt(4)
            address_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run1 = address_para.add_run("OFFICE NO. 12, 2ND FLOOR, NO.40 DLF INDUSTRIAL AREA, KIRTI NAGAR, DELHI 110015.\n")
            set_font(run1, font_name, 9, bold=True, color=text_color_rgb)
            
            run2 = address_para.add_run("Contact No. : ")
            set_font(run2, font_name, 9, bold=True, color=text_color_rgb)
            run3 = address_para.add_run("+91 99871 59986")
            set_font(run3, font_name, 9, bold=True, color=RGBColor(37, 99, 235))
            run4 = address_para.add_run(" | Email : ")
            set_font(run4, font_name, 9, bold=True, color=text_color_rgb)
            run5 = address_para.add_run("info@sgwebapp.com")
            set_font(run5, font_name, 9, bold=True, color=RGBColor(37, 99, 235))
            run6 = address_para.add_run(" | Website : ")
            set_font(run6, font_name, 9, bold=True, color=text_color_rgb)
            run7 = address_para.add_run("www.sgwebapp.com")
            set_font(run7, font_name, 9, bold=True, color=RGBColor(37, 99, 235))
            
            add_horizontal_line(doc, size='16', color='000000', full_bleed=True, left_margin_pt=ml, right_margin_pt=mr) 

        title_p = doc.add_paragraph()
        title_p.paragraph_format.space_before = Pt(12) 
        title_p.paragraph_format.space_after = Pt(24) 
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(doc_title.upper())
        set_font(title_run, font_name, int(font_size) + 4, bold=True, color=text_color_rgb) 

        for idx, section in enumerate(sections):
            heading_title = section.get('heading', 'Section')
            
            head_p = doc.add_paragraph()
            head_p.paragraph_format.space_before = Pt(18) if idx > 0 else Pt(4)
            head_p.paragraph_format.space_after = Pt(6)
            head_run = head_p.add_run(f"{idx + 1}. {heading_title}")
            set_font(head_run, font_name, int(font_size) + 2, bold=True, color=text_color_rgb)
            
            content_str = section.get('content', '')
            format_markdown_to_docx(doc, content_str, font_name, font_size, text_color_rgb)

        stamp_path = os.path.join(assets_dir, "stamp.png")
        if os.path.exists(stamp_path):
            add_horizontal_line(doc, size='4', full_bleed=False)
            stamp_p = doc.add_paragraph()
            stamp_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            stamp_p.add_run().add_picture(stamp_path, width=Inches(1.8))

        # --- SAVE MASTER DOCX ---
        docx_filename = os.path.join(script_dir, f"{file_id}.docx")
        doc.save(docx_filename)

        # --- 3. LINUX-COMPATIBLE PDF EXPORT ---
        if format_type == 'pdf':
            convert_to_pdf_linux(docx_filename, script_dir)
            print("SUCCESS_PDF")
        else:
            print("SUCCESS_DOCX")
            
    except Exception as e:
        print(f"CRITICAL_ERROR:{str(e)}", file=sys.stderr)
        traceback.print_exc(file=sys.stderr)
        sys.exit(1)

if __name__ == "__main__":
    generate_document()